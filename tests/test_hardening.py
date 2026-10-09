"""Hostile or malformed input must be refused cleanly, never crash the server or reach Excel
as a formula."""
from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest
from fastapi.testclient import TestClient

from hazard_ext.engine.core import compute
from hazard_ext.engine.params import Params
from hazard_ext.engine.parse import COLUMNS, ParseError, parse_zip
from hazard_ext.export.formula_xlsx import build_formula_workbook
from hazard_ext.export.tables import summary_workbook
from hazard_ext.export.values_xlsx import build_values_workbook
from hazard_ext.web.config import Settings
from hazard_ext.web.main import create_app

from .conftest import load_zip

HEADER = ",".join(COLUMNS)
H = {"X-Requested-With": "hazard-ext"}


def _zip(rows: list[str], debug: str | None = '{"Category1": "44"}') -> io.BytesIO:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("lgd_recovery.csv", HEADER + "\n" + "\n".join(rows) + "\n")
        if debug is not None:
            zf.writestr("debug.json", debug)
    buf.seek(0)
    return buf


GOOD = "Lifetime,1,0,100,1,100,90,0.1,1,0.98,0.098,0.098,0.902"


def test_minimal_valid_zip_parses():
    data = parse_zip(_zip([GOOD]))
    assert data.category == "44" and data.triangles("Lifetime").n == 1


@pytest.mark.parametrize("row,msg", [
    ("Lifetime,5000000,0,100,1,100,90,0.1,1,0.98,0.098,0.098,0.902", "TermStep must be whole numbers"),
    ("Lifetime,1,0,100,9999999,100,90,0.1,1,0.98,0.098,0.098,0.902", "BucketIndex must be whole numbers"),
    ("Lifetime,0,0,100,1,100,90,0.1,1,0.98,0.098,0.098,0.902", "TermStep must be whole numbers"),
    ("Lifetime,1.5,0,100,2,100,90,0.1,1,0.98,0.098,0.098,0.902", "TermStep must be whole numbers"),
    ("Lifetime,1,0,,1,100,90,0.1,1,0.98,0.098,0.098,0.902", "blank or non-numeric"),
    ("Lifetime,1,0,abc,1,100,90,0.1,1,0.98,0.098,0.098,0.902", "non-numeric"),
])
def test_rows_that_would_blow_up_the_engine_are_refused(row, msg):
    with pytest.raises(ParseError, match=msg):
        parse_zip(_zip([GOOD, row]))


def test_debug_json_must_be_an_object():
    with pytest.raises(ParseError, match="JSON object"):
        parse_zip(_zip([GOOD], debug="[1, 2]"))
    with pytest.raises(ParseError, match="not valid JSON"):
        parse_zip(_zip([GOOD], debug="{nope"))
    assert parse_zip(_zip([GOOD], debug=None)).category == ""


def test_names_that_look_like_formulas_stay_text_in_every_workbook():
    evil = '=HYPERLINK("http://example.com","x")'
    data = load_zip("44")
    res = compute(data, Params(method=3, target_ts=120, max_bucket=240))

    wb = openpyxl.load_workbook(io.BytesIO(build_values_workbook(res, evil, evil, evil)))
    assert wb["Charts"]["A1"].data_type == "s" and wb["Charts"]["A1"].value.startswith("=HYPERLINK")
    assert wb["Summary"]["A2"].data_type == "s"

    wb = openpyxl.load_workbook(io.BytesIO(build_formula_workbook(res, data, evil, evil)))
    assert wb["Charts"]["A1"].data_type == "s"
    assert wb["README"]["B1"].data_type == "s"

    matrix = {"datasets": [{"id": 1, "name": evil, "category": evil}], "scenarios": [{"id": 1, "name": evil}],
              "cells": [{"dataset_id": 1, "scenario_id": 1, "status": "error", "error": evil, "summary": {}}]}
    ws = openpyxl.load_workbook(io.BytesIO(summary_workbook(evil, matrix)))["Summary"]
    for col in ("A", "B", "C", "U"):
        assert ws[f"{col}4"].data_type == "s", col


def test_oversized_upload_is_refused_before_it_is_read(tmp_path):
    settings = Settings(_env_file=None, database_url=f"sqlite:///{(tmp_path / 'a.db').as_posix()}",
                        local_data_dir=str(tmp_path / "blobs"), admin_email="a@example.com",
                        admin_password="long-password-1", max_upload_mb=1)
    with TestClient(create_app(settings)) as c:
        assert c.post("/api/auth/login", json={"email": "a@example.com", "password": "long-password-1"}, headers=H).status_code == 200
        pid = c.post("/api/projects", json={"name": "P"}, headers=H).json()["id"]
        big = ("files", ("big.zip", b"\0" * (3 * 1024 * 1024), "application/zip"))
        r = c.post(f"/api/projects/{pid}/datasets", files=[big], headers=H)
        assert r.status_code == 413 and "1 MB" in r.json()["detail"]
        small = ("files", ("ok.zip", _zip([GOOD]).getvalue(), "application/zip"))
        out = c.post(f"/api/projects/{pid}/datasets", files=[small], headers=H).json()
        assert out[0]["ok"] is True and out[0]["dataset"]["name"] == "VB44"


def test_sign_in_is_throttled_per_account(tmp_path):
    settings = Settings(_env_file=None, database_url=f"sqlite:///{(tmp_path / 'b.db').as_posix()}",
                        local_data_dir=str(tmp_path / "blobs"), admin_email="a@example.com",
                        admin_password="long-password-1")
    with TestClient(create_app(settings)) as c:
        for _ in range(8):
            assert c.post("/api/auth/login", json={"email": "a@example.com", "password": "wrong"}, headers=H).status_code == 401
        blocked = c.post("/api/auth/login", json={"email": "a@example.com", "password": "long-password-1"}, headers=H)
        assert blocked.status_code == 429
