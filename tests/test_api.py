"""API behaviour on SQLite with local storage. Tests in this module run in order and share
one application, as a user session would."""
from __future__ import annotations

import io

import numpy as np
import openpyxl
import pytest
from fastapi.testclient import TestClient

from sqlalchemy import select

from hazard_ext.engine.core import compute
from hazard_ext.engine.params import Params
from hazard_ext.engine.parse import RecoveryData
from hazard_ext.web.config import Settings
from hazard_ext.web.main import create_app
from hazard_ext.web.models import Dataset, Result, Scenario, ScenarioOverride
from hazard_ext.web.runner import migrate_three_methods

from .conftest import load_zip, zip_path

H = {"X-Requested-With": "hazard-ext"}
ADMIN = ("admin@example.com", "admin-password-1")
EDITOR = ("editor@example.com", "editor-password-1")
VIEWER = ("viewer@example.com", "viewer-password-1")
OUTSIDER = ("outsider@example.com", "outsider-password-1")


class World:
    """Shared state across the ordered tests."""


@pytest.fixture(scope="module")
def world(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("app")
    settings = Settings(
        _env_file=None,
        database_url=f"sqlite:///{(tmp / 'app.db').as_posix()}",
        local_data_dir=str(tmp / "blobs"),
        session_secret="test-secret",
        admin_email=ADMIN[0], admin_password=ADMIN[1],
    )
    app = create_app(settings)
    w = World()
    w.tmp = tmp
    with TestClient(app) as admin:
        w.app = app
        w.admin = admin
        w.anon = TestClient(app)
        w.editor = TestClient(app)
        w.viewer = TestClient(app)
        w.outsider = TestClient(app)
        yield w


def login(client, who):
    r = client.post("/api/auth/login", json={"email": who[0], "password": who[1]}, headers=H)
    assert r.status_code == 200, r.text
    return r.json()


def upload(client, pid, names):
    files = []
    for name in names:
        if isinstance(name, tuple):
            files.append(("files", (name[0], name[1], "application/zip")))
        else:
            files.append(("files", (f"debug ({name}).zip", zip_path(name).read_bytes(), "application/zip")))
    return client.post(f"/api/projects/{pid}/datasets", files=files, headers=H)


# ----------------------------------------------------------------- sign-in
def test_health_and_static_page(world):
    assert world.anon.get("/api/health").json() == {"ok": True}
    page = world.anon.get("/")
    assert page.status_code == 200 and "<html" in page.text.lower()


def test_api_needs_sign_in(world):
    assert world.anon.get("/api/projects").status_code == 401
    assert world.anon.get("/api/auth/me").status_code == 401


def test_wrong_password_is_rejected(world):
    r = world.anon.post("/api/auth/login", json={"email": ADMIN[0], "password": "nope-nope-nope"}, headers=H)
    assert r.status_code == 401
    r = world.anon.post("/api/auth/login", json={"email": "ghost@example.com", "password": "x"}, headers=H)
    assert r.status_code == 401


def test_mutation_without_header_is_rejected(world):
    r = world.anon.post("/api/auth/login", json={"email": ADMIN[0], "password": ADMIN[1]})
    assert r.status_code == 403


def test_admin_signs_in_and_creates_users(world):
    me = login(world.admin, ADMIN)
    assert me["is_admin"] is True
    for who in (EDITOR, VIEWER, OUTSIDER):
        r = world.admin.post("/api/admin/users", json={"email": who[0], "password": who[1], "name": who[0]}, headers=H)
        assert r.status_code == 201, r.text
    assert world.admin.post("/api/admin/users", json={"email": EDITOR[0], "password": "another-pass-1"},
                            headers=H).status_code == 409
    assert world.admin.post("/api/admin/users", json={"email": "short@example.com", "password": "short"},
                            headers=H).status_code == 422
    assert len(world.admin.get("/api/admin/users").json()) == 4


def test_non_admin_cannot_manage_users(world):
    login(world.editor, EDITOR)
    assert world.editor.get("/api/admin/users").status_code == 403
    assert world.editor.post("/api/admin/users", json={"email": "x@example.com", "password": "long-password-1"},
                             headers=H).status_code == 403


def test_last_admin_cannot_be_removed(world):
    me = world.admin.get("/api/auth/me").json()
    r = world.admin.patch(f"/api/admin/users/{me['id']}", json={"is_admin": False}, headers=H)
    assert r.status_code == 400


# ---------------------------------------------------------------- projects
def test_project_membership_and_isolation(world):
    r = world.editor.post("/api/projects", json={"name": "Nutun July 2026"}, headers=H)
    assert r.status_code == 201
    world.pid = pid = r.json()["id"]
    login(world.viewer, VIEWER)
    login(world.outsider, OUTSIDER)

    assert world.outsider.get(f"/api/projects/{pid}").status_code == 404     # not a member: invisible
    assert world.outsider.get("/api/projects").json() == []
    assert world.viewer.get(f"/api/projects/{pid}").status_code == 404

    r = world.editor.put(f"/api/projects/{pid}/members", json={"email": VIEWER[0], "role": "viewer"}, headers=H)
    assert r.status_code == 200 and {m["email"]: m["role"] for m in r.json()} == {
        EDITOR[0]: "owner", VIEWER[0]: "viewer"}
    assert world.viewer.get(f"/api/projects/{pid}").json()["role"] == "viewer"
    assert world.editor.put(f"/api/projects/{pid}/members", json={"email": "nobody@example.com", "role": "viewer"},
                            headers=H).status_code == 404
    # the only owner cannot be demoted or removed
    editor_id = world.editor.get("/api/auth/me").json()["id"]
    assert world.editor.put(f"/api/projects/{pid}/members", json={"email": EDITOR[0], "role": "viewer"},
                            headers=H).status_code == 400
    assert world.editor.delete(f"/api/projects/{pid}/members/{editor_id}", headers=H).status_code == 400
    # viewers cannot manage members
    assert world.viewer.put(f"/api/projects/{pid}/members", json={"email": OUTSIDER[0], "role": "viewer"},
                            headers=H).status_code == 403


# ------------------------------------------------------------------ upload
def test_upload_several_zips_each_judged_on_its_own(world):
    r = upload(world.editor, world.pid, ["44", ("broken.zip", b"not a zip at all"), "ALL"])
    assert r.status_code == 200
    out = r.json()
    assert [o["ok"] for o in out] == [True, False, True]
    assert "valid zip" in out[1]["error"]
    d44, dall = out[0]["dataset"], out[2]["dataset"]
    assert (d44["name"], d44["category"], d44["profile"]["n"]) == ("VB44", "44", 97)
    assert (dall["name"], dall["category"], dall["profile"]["n"]) == ("VBALL", "ALL", 329)
    assert d44["profile"]["implied_rate"] == pytest.approx(0.1771, abs=1e-6)
    world.d44, world.dall = d44["id"], dall["id"]
    assert len(list((world.tmp / "blobs" / f"p{world.pid}").glob("*.npz"))) == 2


def test_duplicate_zip_is_rejected(world):
    out = upload(world.editor, world.pid, ["44"]).json()
    assert out[0]["ok"] is False and "already in the project" in out[0]["error"]


def test_viewer_is_read_only(world):
    pid = world.pid
    assert upload(world.viewer, pid, ["22"]).status_code == 403
    assert world.viewer.post(f"/api/projects/{pid}/scenarios", json={"name": "x"}, headers=H).status_code == 403
    assert world.viewer.delete(f"/api/projects/{pid}/datasets/{world.d44}", headers=H).status_code == 403
    assert len(world.viewer.get(f"/api/projects/{pid}").json()["datasets"]) == 2


# --------------------------------------------------------------- scenarios
def test_scenario_defaults_and_validation(world):
    pid = world.pid
    r = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Base", "params": {"method": 3, "target_ts": 360, "max_bucket": 480}}, headers=H)
    assert r.status_code == 201, r.text
    s = r.json()
    world.sid = s["id"]
    assert s["params"]["method"] == 3 and s["params"]["target_ts"] == 360 and s["params"]["window"] == 12
    assert s["params"]["vintage_start"] is None and s["params"]["vintage_years"] is None and "client_cohort" not in s["params"]
    plain = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Plain"}, headers=H)
    assert plain.status_code == 201 and plain.json()["params"]["method"] == 1
    # the retired reference-curve parameter is dropped, so a stale browser or an old saved set still loads
    old = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Old", "params": {"client_cohort": "44", "window": 6}}, headers=H)
    assert old.status_code == 201 and "client_cohort" not in old.json()["params"] and old.json()["params"]["window"] == 6
    assert world.editor.delete(f"/api/projects/{pid}/scenarios/{old.json()['id']}", headers=H).status_code == 200
    # the two vintage parameters are exclusive, and the month is normalised
    both = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Both", "params": {"vintage_start": "2016-08", "vintage_years": 10}}, headers=H)
    assert both.status_code == 422 and "not both" in both.text
    v = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Vint", "params": {"vintage_start": "2016-08-31"}}, headers=H)
    assert v.status_code == 201 and v.json()["params"]["vintage_start"] == "2016-08"
    assert world.editor.delete(f"/api/projects/{pid}/scenarios/{v.json()['id']}", headers=H).status_code == 200
    assert world.editor.delete(f"/api/projects/{pid}/scenarios/{plain.json()['id']}", headers=H).status_code in (200, 204)
    assert world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Base"}, headers=H).status_code == 409
    bad = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Bad", "params": {"windw": 3}}, headers=H)
    assert bad.status_code == 422 and "windw" in bad.text
    bad = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Bad", "params": {"method": 7}}, headers=H)
    assert bad.status_code == 422


def test_run_records_a_failure_per_zip_without_stopping_the_others(world):
    pid, sid = world.pid, world.sid
    assert world.viewer.post(f"/api/projects/{pid}/scenarios/{sid}/run", json={}, headers=H).status_code == 403
    # a vintage start after VB44's last vintage fails for that zip only
    world.editor.put(f"/api/projects/{pid}/scenarios/{sid}/overrides/{world.d44}", json={"params": {"vintage_start": "2030-01"}}, headers=H)
    out = world.editor.post(f"/api/projects/{pid}/scenarios/{sid}/run", json={}, headers=H).json()
    by = {o["dataset"]: o for o in out}
    assert by["VBALL"]["status"] == "ok" and by["VBALL"]["summary"]["method_label"] == "Log-normal"
    assert by["VB44"]["status"] == "error" and "after the last vintage" in by["VB44"]["error"]
    world.editor.put(f"/api/projects/{pid}/scenarios/{sid}/overrides/{world.d44}", json={"params": {}}, headers=H)
    out = world.editor.post(f"/api/projects/{pid}/scenarios/{sid}/run", json={}, headers=H).json()
    assert {o["dataset"]: o["status"] for o in out} == {"VB44": "ok", "VBALL": "ok"}


def test_result_equals_the_engine(world):
    r = world.viewer.get(f"/api/projects/{world.pid}/results/{world.sid}/{world.d44}").json()
    expect = compute(load_zip("44"), Params(method=3, target_ts=360, max_bucket=480))
    assert r["status"] == "ok" and r["stale"] is False and r["legacy"] is False
    assert r["payload"]["config"]["mu"] == pytest.approx(expect.config["mu"]) and r["summary"]["sigma"] == pytest.approx(expect.config["sigma"])
    assert r["summary"]["vintages"] == "" and r["payload"]["config"]["vintage_first"] == "2018-06"
    assert r["summary"]["lgd_selected"] == pytest.approx(expect.averages["lgd_selected"], abs=1e-12)
    np.testing.assert_allclose(r["payload"]["results"]["lgd_selected"], expect.results["lgd_selected"], atol=1e-12)
    assert len(r["payload"]["lgd_ts"]["lgd_final"]) == 360
    assert r["effective_params"]["max_bucket"] == 480


def test_override_takes_precedence_for_one_zip_only(world):
    pid, sid = world.pid, world.sid
    r = world.editor.put(f"/api/projects/{pid}/scenarios/{sid}/overrides/{world.dall}",
                         json={"params": {"method": 1, "min_exposure_mode": "pct", "min_exposure": 0.5}}, headers=H)
    assert r.status_code == 200 and r.json()["overrides"] == {
        str(world.dall): {"method": 1, "min_exposure_mode": "pct", "min_exposure": 0.5}}
    bad = world.editor.put(f"/api/projects/{pid}/scenarios/{sid}/overrides/{world.dall}",
                           json={"params": {"nope": 1}}, headers=H)
    assert bad.status_code == 422
    out = world.editor.post(f"/api/projects/{pid}/scenarios/{sid}/run", json={"dataset_id": world.dall}, headers=H).json()
    assert out[0]["status"] == "ok" and out[0]["summary"]["method"] == 1
    r44 = world.editor.get(f"/api/projects/{pid}/results/{sid}/{world.d44}").json()
    rall = world.editor.get(f"/api/projects/{pid}/results/{sid}/{world.dall}").json()
    assert r44["effective_params"]["method"] == 3 and rall["effective_params"]["method"] == 1
    assert rall["payload"]["config"]["min_exposure_abs"] == pytest.approx(0.005 * rall["payload"]["config"]["opening_exposure"])
    assert rall["payload"]["averages"]["lgd_logn"] is not None     # every shape is computed whatever the method


def test_matrix(world):
    m = world.viewer.get(f"/api/projects/{world.pid}/matrix").json()
    assert [d["name"] for d in m["datasets"]] == ["VB44", "VBALL"]
    cells = {(c["scenario_id"], c["dataset_id"]): c for c in m["cells"]}
    assert cells[(world.sid, world.d44)]["status"] == "ok" and not cells[(world.sid, world.d44)]["has_override"]
    assert cells[(world.sid, world.dall)]["has_override"] is True


def test_series_gives_final_lgd_for_every_computed_result(world):
    rows = world.viewer.get(f"/api/projects/{world.pid}/series").json()
    assert {(r["scenario_id"], r["dataset_id"]) for r in rows} == {(world.sid, world.d44), (world.sid, world.dall)}
    r44 = next(r for r in rows if r["dataset_id"] == world.d44)
    assert len(r44["lgd_final"]) == 360 and r44["ts"][0] == 1 and r44["stale"] is False
    assert r44["lgd_selected_avg"] == pytest.approx(
        world.viewer.get(f"/api/projects/{world.pid}/results/{world.sid}/{world.d44}").json()["summary"]["lgd_selected"])
    assert world.outsider.get(f"/api/projects/{world.pid}/series").status_code == 404


def test_editing_a_scenario_marks_results_stale_until_rerun(world):
    pid, sid = world.pid, world.sid
    s = world.editor.get(f"/api/projects/{pid}/scenarios/{sid}").json()
    same = world.editor.put(f"/api/projects/{pid}/scenarios/{sid}", json={"params": s["params"]}, headers=H)
    assert same.status_code == 200
    assert world.editor.get(f"/api/projects/{pid}/results/{sid}/{world.d44}").json()["stale"] is False
    params = {**s["params"], "window": 24}
    assert world.editor.put(f"/api/projects/{pid}/scenarios/{sid}", json={"params": params}, headers=H).status_code == 200
    assert world.editor.get(f"/api/projects/{pid}/results/{sid}/{world.d44}").json()["stale"] is True
    world.editor.post(f"/api/projects/{pid}/scenarios/{sid}/run", json={}, headers=H)
    r = world.editor.get(f"/api/projects/{pid}/results/{sid}/{world.d44}").json()
    assert r["stale"] is False and r["effective_params"]["window"] == 24


def test_clone_copies_parameters_and_overrides(world):
    pid = world.pid
    r = world.editor.post(f"/api/projects/{pid}/scenarios/{world.sid}/clone", json={"name": "Exponential"}, headers=H)
    assert r.status_code == 201
    c = r.json()
    world.sid2 = c["id"]
    assert c["params"]["window"] == 24 and str(world.dall) in c["overrides"]
    params = {**c["params"], "method": 1}
    world.editor.put(f"/api/projects/{pid}/scenarios/{c['id']}", json={"params": params}, headers=H)
    assert world.editor.post(f"/api/projects/{pid}/run-all", headers=H).json() == {"ok": 4, "errors": 0}


def test_stale_result_cannot_be_exported_to_excel_until_rerun(world):
    pid, sid, did = world.pid, world.sid, world.d44
    s = world.editor.get(f"/api/projects/{pid}/scenarios/{sid}").json()
    world.editor.put(f"/api/projects/{pid}/scenarios/{sid}", json={"params": {**s["params"], "window": 18}}, headers=H)
    base = f"/api/projects/{pid}/results/{sid}/{did}/export"
    r = world.viewer.get(base + "?kind=values")
    assert r.status_code == 409 and "Run the scenario again" in r.json()["detail"]
    assert world.viewer.get(base + "?kind=csv&table=results").status_code == 200     # the stored tables still download
    world.editor.put(f"/api/projects/{pid}/scenarios/{sid}", json={"params": s["params"]}, headers=H)
    world.editor.post(f"/api/projects/{pid}/scenarios/{sid}/run", json={}, headers=H)
    assert world.viewer.get(base + "?kind=values").status_code == 200


def test_curve_data_for_a_termstep(world):
    r = world.viewer.get(f"/api/projects/{world.pid}/results/{world.sid}/{world.d44}/curve?ts=48").json()
    assert r["bucket"][0] == 48 and r["bucket"][-1] == 480
    assert len(r["logn"]) == len(r["bucket"]) and r["observed"][0] is not None and "client" not in r
    assert r["observed"][-1] is None
    bad = world.viewer.get(f"/api/projects/{world.pid}/results/{world.sid}/{world.d44}/curve?ts=999")
    assert bad.status_code == 400


# ------------------------------------------------------------------ curves
def test_reference_shape_curves_are_refused_and_applied_curves_still_work(world):
    pid = world.pid
    text = "t,ALL\n1,0.01\n2,0.02\n"
    assert world.viewer.post(f"/api/projects/{pid}/curves", files={"file": ("all.csv", text, "text/csv")},
                             data={"kind": "applied"}, headers=H).status_code == 403
    before = world.editor.get(f"/api/projects/{pid}/results/{world.sid}/{world.d44}").json()
    r = world.editor.post(f"/api/projects/{pid}/curves", files={"file": ("all.csv", text, "text/csv")}, data={"kind": "shape"}, headers=H)
    assert r.status_code == 422 and "no longer used" in r.json()["detail"]
    assert world.editor.get(f"/api/projects/{pid}/curves").json() == []
    assert world.editor.get(f"/api/projects/{pid}/curves/ALL?kind=shape").status_code == 404
    # the default kind is applied, and uploading one never marks results out of date
    r = world.editor.post(f"/api/projects/{pid}/curves", files={"file": ("all.csv", text, "text/csv")}, headers=H)
    assert r.status_code == 200 and [c["kind"] for c in r.json()] == ["applied"]
    assert world.editor.get(f"/api/projects/{pid}/results/{world.sid}/{world.d44}").json()["stale"] is False
    assert world.editor.get(f"/api/projects/{pid}/results/{world.sid}/{world.d44}").json()["computed_at"] == before["computed_at"]
    world.editor.delete(f"/api/projects/{pid}/curves/{r.json()[0]['id']}", headers=H)
    bad = world.editor.post(f"/api/projects/{pid}/curves", files={"file": ("x.csv", "t,a\n1,1\n5,2\n", "text/csv")}, headers=H)
    assert bad.status_code == 422
    # remove ALL's override so both zips run on the scenario's log-normal method
    world.editor.put(f"/api/projects/{pid}/scenarios/{world.sid}/overrides/{world.dall}", json={"params": {}}, headers=H)
    out = world.editor.post(f"/api/projects/{pid}/scenarios/{world.sid}/run", json={}, headers=H).json()
    assert {o["dataset"]: o["status"] for o in out} == {"VB44": "ok", "VBALL": "ok"}


def test_applied_curves_compare_without_touching_the_calculation(world):
    pid, sid, d44 = world.pid, world.sid, world.d44
    before = world.editor.get(f"/api/projects/{pid}/results/{sid}/{d44}").json()
    # an outstanding-basis curve for cohort 44: 5% of the outstanding balance every month for 120 months
    text = "t,44\n" + "\n".join(f"{t},0.05" for t in range(1, 121))
    r = world.editor.post(f"/api/projects/{pid}/curves", files={"file": ("applied.csv", text, "text/csv")},
                          data={"kind": "applied", "basis": "outstanding"}, headers=H)
    assert r.status_code == 200, r.text
    rows = {(c["label"], c["kind"]): c for c in r.json()}
    assert rows[("44", "applied")]["basis"] == "outstanding" and len(rows) == 1
    assert abs(rows[("44", "applied")]["total"] - (1 - 0.95 ** 120)) < 1e-9       # converted to the face basis
    # results are not marked out of date: applied curves are comparison only
    after = world.editor.get(f"/api/projects/{pid}/results/{sid}/{d44}").json()
    assert after["stale"] is False and after["computed_at"] == before["computed_at"]

    curve = world.viewer.get(f"/api/projects/{pid}/curves/44?kind=applied").json()
    assert curve["values"][0] == pytest.approx(0.05) and curve["values"][1] == pytest.approx(0.05 * 0.95)
    assert world.viewer.get(f"/api/projects/{pid}/curves/44").json()["kind"] == "applied"

    c = world.viewer.get(f"/api/projects/{pid}/results/{sid}/{d44}/curve?ts=3").json()
    assert c["applied_label"] == "44" and c["applied_basis"] == "outstanding"
    assert c["applied"][0] == pytest.approx(0.05 * 0.95 ** 2 / 0.95 ** 2)         # rolled forward to TermStep 3
    assert c["applied"][-1] is None                                                 # beyond the 120 uploaded months
    a = world.viewer.get(f"/api/projects/{pid}/results/{sid}/{d44}/applied").json()
    assert a["available"] is True and len(a["lgd"]) == 360 and 0 < a["lgd"][0] < 1
    # no applied curve for the ALL zip
    assert world.viewer.get(f"/api/projects/{pid}/results/{sid}/{world.dall}/applied").json()["available"] is False
    # comparison on both bases
    cf = world.viewer.get(f"/api/projects/{pid}/results/{sid}/{d44}/compare_curves?ts=1&basis=face").json()
    co = world.viewer.get(f"/api/projects/{pid}/results/{sid}/{d44}/compare_curves?ts=1&basis=outstanding").json()
    assert cf["applied_label"] == "44" and cf["selected"] == "logn" and cf["series"]["applied"][0] == pytest.approx(0.05)
    cp = {p_["shape"]: p_ for p_ in cf["curve_params"]}
    assert cp["logn"]["selected"] and cp["logn"]["params"][0]["name"] == "μ" and cp["logn"]["scale_ts"] > 0
    assert cp["exp"]["params"][0]["fitted"] == pytest.approx(0.0528, abs=1e-3) and cp["exp"]["formula"].startswith("s(b)")
    assert set(cf["series"]) == {"observed", "exp", "power", "logn", "applied"} and "reference_label" not in cf
    assert co["series"]["applied"][0] == pytest.approx(0.05) and co["series"]["applied"][1] == pytest.approx(0.05)   # back on its own basis
    assert cf["series"]["exp"][0] == pytest.approx(co["series"]["exp"][0])                                        # first step is the same on both bases
    assert cf["cumulative"][0]["months"] == 12 and cf["cumulative"][0]["client"] == pytest.approx(1 - 0.95 ** 12)
    assert cf["cumulative"][0]["difference"] == pytest.approx(cf["cumulative"][0]["ours"] - cf["cumulative"][0]["client"])
    assert world.viewer.get(f"/api/projects/{pid}/results/{sid}/{d44}/compare_curves?basis=nope").status_code == 422
    # validation
    assert world.editor.post(f"/api/projects/{pid}/curves", files={"file": ("x.csv", "t,44\n1,1.5\n", "text/csv")},
                             data={"kind": "applied", "basis": "outstanding"}, headers=H).status_code == 422
    r = world.editor.post(f"/api/projects/{pid}/curves", files={"file": ("x.csv", "t,44\n1,0.6\n2,0.6\n", "text/csv")},
                          data={"kind": "applied", "basis": "face"}, headers=H)
    assert r.status_code == 422 and "add up to more" in r.json()["detail"]        # 120% of face is impossible
    # removing it does not mark results out of date either
    world.editor.delete(f"/api/projects/{pid}/curves/{rows[('44', 'applied')]['id']}", headers=H)
    assert world.editor.get(f"/api/projects/{pid}/results/{sid}/{d44}").json()["stale"] is False
    assert world.viewer.get(f"/api/projects/{pid}/results/{sid}/{d44}/applied").json()["available"] is False


# ----------------------------------------------------------------- exports
def test_csv_export(world):
    r = world.viewer.get(f"/api/projects/{world.pid}/results/{world.sid}/{world.d44}/export?kind=csv&table=lgd_ts")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    lines = r.text.strip().split("\n")
    assert lines[0].startswith("TermStep,") and len(lines) == 361
    assert world.viewer.get(f"/api/projects/{world.pid}/results/{world.sid}/{world.d44}/export?kind=csv&table=zzz").status_code == 422


def _engine_44():
    return compute(load_zip("44"), Params(method=3, target_ts=360, max_bucket=480, window=24))


def test_values_workbook(world):
    r = world.viewer.get(f"/api/projects/{world.pid}/results/{world.sid}/{world.d44}/export?kind=values")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    assert wb.sheetnames == ["Summary", "Results", "LGD_TS", "Tail_Fit", "Chart_Data", "Charts", "Hazard_Obs",
                             "Exposure_Obs", "Ext_Exp", "Ext_Power", "Ext_LogN"]
    e = _engine_44()
    ws = wb["Results"]
    heads = [c.value for c in ws[1]]
    col = heads.index("LGD – SELECTED") + 1
    got = [ws.cell(row=1 + ts, column=col).value for ts in range(1, 98)]
    np.testing.assert_allclose(got, e.results["lgd_selected"], atol=1e-12)
    assert "reference" not in " ".join(str(c.value) for c in wb["Results"][1]).lower()
    ws = wb["Ext_LogN"]
    assert ws.cell(row=4, column=1).value == 1 and ws.max_column == 481
    assert ws.cell(row=4, column=1 + 200).value == pytest.approx(e.ext[2, 1, 200], abs=1e-15)
    assert ws.cell(row=4 + 49, column=1 + 10).value is None           # b < ts stays blank
    assert len(wb["Charts"]._charts) == 6


def test_formula_workbook(world):
    r = world.viewer.get(f"/api/projects/{world.pid}/results/{world.sid}/{world.d44}/export?kind=formula")
    assert r.status_code == 200, r.text
    e = _engine_44()
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=False)
    assert wb.sheetnames == ["README", "Config", "Results", "LGD_300", "Charts", "Hazard_Obs", "Tail_Fit",
                             "Ext_Exp", "Ext_Power", "Ext_LogN", "Raw_Index", "Raw_Debug"]
    names = set(wb.defined_names.keys())
    assert {"BaseTS", "MuOvr", "SigOvr", "Mu", "Sig", "EventType", "FitStart", "Floor", "Gam", "GamOvr", "Horizon", "Horizon2", "Lam",
            "LamOvr", "LastTS", "MaxBucket", "Method", "MinExp", "Rate", "RefLastCred", "RefTS", "v", "WinW"} <= names
    assert "RefCurve" not in names
    cfg = wb["Config"]
    assert cfg["B6"].value == 480 and cfg["B8"].value == 24 and cfg["B11"].value == 3 and cfg["B12"].value is None
    assert cfg["B22"].value == "=(1+Rate)^(-1/12)"
    assert cfg["B27"].value == '=IF(LamOvr="",B25,LamOvr)'
    assert cfg["B32"].value == "=IF(B35<0,SQRT(-1/(2*B35)),NA())" and cfg["B33"].value.startswith('=IF(MuOvr<>"",MuOvr,')
    assert wb["Results"]["I10"].value == "=CHOOSE(Method,F10,G10,H10)"
    assert wb["Ext_Exp"]["C6"].value == "=IF(C$5<$A6,0,IF(C$5<=$B6,Hazard_Obs!C$6,MAX(Floor,Tail_Fit!$H$20*Tail_Fit!C$6)))"
    assert wb["Tail_Fit"]["D20"].value == "=MAX(A20,C20-WinW+1)"
    assert wb["Tail_Fit"]["C8"].value == "=EXP(-((LN(C$5)-Mu)^2)/(2*Sig^2))/C$5"
    assert wb["Tail_Fit"]["C12"].value == '=IF(ISNUMBER(C$10),C$10+C$11,"")'
    assert wb["Raw_Debug"].max_row == 3 + 97 * 98 // 2            # one EventType block
    assert len(wb["Charts"]._charts) == 6

    # cached values are the engine's values
    wv = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    res = wv["Results"]
    np.testing.assert_allclose([res.cell(row=9 + ts, column=9).value for ts in range(1, 98)],
                               e.results["lgd_selected"], atol=1e-12)
    assert res["G6"].value == pytest.approx(e.averages["lgd_selected"], abs=1e-12)
    lgd = wv["LGD_300"]
    assert lgd.max_row == 12 + 360
    np.testing.assert_allclose([lgd.cell(row=12 + ts, column=10).value for ts in range(1, 361)],
                               e.lgd_ts["lgd_final"], atol=1e-12)
    assert lgd.cell(row=12 + 200, column=15).value == "derived"
    assert wv["Config"]["B25"].value == pytest.approx(e.config["lam_fit"], abs=1e-12)
    assert wv["Config"]["B31"].value == pytest.approx(e.config["mu_fit"], abs=1e-12)
    assert wv["Config"]["B34"].value == pytest.approx(e.config["sigma"], abs=1e-12)
    assert wv["Hazard_Obs"]["C6"].value == pytest.approx(e.R[1, 1], abs=1e-15)
    assert wv["Raw_Index"]["B6"].value == 97 + 1 and wv["Raw_Index"]["C6"].value == 96
    assert wv["Ext_LogN"].cell(row=5 + 1, column=2 + 300).value == pytest.approx(e.ext[2, 1, 300], abs=1e-15)
    text = " ".join(str(c.value) for row in wv["README"].iter_rows() for c in row if c.value) + " ".join(wv.sheetnames)
    assert "reference curve" not in text.lower() and "Ref_Curve" not in text


def test_summary_workbook(world):
    r = world.viewer.get(f"/api/projects/{world.pid}/export/summary")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames[:7] == ["Summary", "LGD_by_TS", "LGD_final_wide", "Scenario_deltas", "Cumulative_recovery",
                                 "Exposure_credibility", "Charts"]
    assert sorted(n for n in wb.sheetnames if n.startswith("M_")) == ["M_VB44_Base", "M_VB44_Exponential", "M_VBALL_Base", "M_VBALL_Exponential"]
    ws = wb["Summary"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=5, max_col=18)]
    assert sorted((r[0], r[2]) for r in rows) == [("VB44", "Base"), ("VB44", "Exponential"),
                                                    ("VBALL", "Base"), ("VBALL", "Exponential")]
    e = _engine_44()
    r44 = next(r for r in rows if r[0] == "VB44" and r[2] == "Base")
    assert r44[17] == pytest.approx(e.averages["lgd_selected"], abs=1e-12)        # LGD SELECTED
    assert r44[4] == "Log-normal" and r44[5] == "All vintages"
    heads = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    assert "LGD log-normal" in heads and "Log-normal μ used" in heads and not any("reference curve" in str(h).lower() for h in heads)
    # marginal sheet: TermStep rows, remaining steps 1..120, own rows equal the extended row
    m = wb["M_VB44_Base"]
    cells = [[c.value for c in row] for row in m.iter_rows(min_row=2, max_row=4, max_col=8)]
    assert cells[0][:5] == ["TermStep", "Source", "LGD final", "Σ steps 1–120", 1]
    assert cells[1][0] == 1 and cells[1][1] == "own row"
    assert cells[1][4] == pytest.approx(e.ext[2, 1, 1], abs=1e-15) and cells[1][5] == pytest.approx(e.ext[2, 1, 2], abs=1e-15)
    assert cells[2][4] == pytest.approx(e.ext[2, 2, 2], abs=1e-15)
    # LGD_by_TS has one row per zip, scenario and TermStep
    assert wb["LGD_by_TS"].max_row == 1 + 4 * 360
    assert wb["Scenario_deltas"]["B4"].value == "Exponential"


def test_summary_workbook_when_the_target_is_below_the_observed_range(world):
    # VBALL has 329 observed TermSteps; a Target of 200 must not break the weighted averages
    pid = world.pid
    s = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Short", "params": {"target_ts": 200, "max_bucket": 320}}, headers=H).json()
    world.editor.put(f"/api/projects/{pid}/scenarios/{s['id']}/overrides/{world.dall}", json={"params": {"method": 1}}, headers=H)
    world.editor.post(f"/api/projects/{pid}/scenarios/{s['id']}/run", json={}, headers=H)
    r = world.viewer.get(f"/api/projects/{pid}/export/summary")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert "M_VBALL_Short" in wb.sheetnames
    world.editor.delete(f"/api/projects/{pid}/scenarios/{s['id']}", headers=H)


def test_curves_workbook_for_one_scenario(world):
    pid, sid = world.pid, world.sid
    r = world.viewer.get(f"/api/projects/{pid}/export/curves?scenario_id={sid}")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames == ["Notes", "Curve_parameters", "Scale_by_TermStep", "LGD_by_TermStep", "Marginal_face",
                             "Marginal_outstanding", "Cumulative_face", "M_VB44", "M_VBALL"]
    e = _engine_44()
    cp = wb["Curve_parameters"]
    heads = [c.value for c in next(cp.iter_rows(min_row=4, max_row=4))]
    row44 = [c.value for c in next(cp.iter_rows(min_row=5, max_row=5))]
    assert row44[0] == "VB44" and row44[heads.index("λ fitted")] == pytest.approx(e.config["lam_fit"], abs=1e-12)
    assert row44[heads.index("μ used")] == pytest.approx(e.config["mu"], abs=1e-12)
    assert row44[heads.index("Scale log-normal, TermStep 1")] == pytest.approx(e.tail_fit["scale_logn"][0], rel=1e-12)
    sc = wb["Scale_by_TermStep"]
    assert [c.value for c in next(sc.iter_rows(min_row=4, max_row=4))][1] == "VB44 – scale exponential"
    ws = wb["LGD_by_TermStep"]
    heads = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    assert heads[:3] == ["TermStep", "VB44 – final LGD", "VBALL – final LGD"]
    assert ws["B5"].value == pytest.approx(e.lgd_ts["lgd_final"][0], abs=1e-12)
    mf = wb["Marginal_face"]
    heads = [c.value for c in next(mf.iter_rows(min_row=4, max_row=4))]
    assert heads[1] == "VB44 – ours (Log-normal)" and not any("reference" in str(h).lower() for h in heads)
    assert mf["B5"].value == pytest.approx(e.ext[2, 1, 1], abs=1e-15)
    mo = wb["Marginal_outstanding"]
    assert mo["B5"].value == pytest.approx(e.ext[2, 1, 1], abs=1e-15)                        # month 1 equal on both bases
    assert mo["B6"].value == pytest.approx(e.ext[2, 1, 2] / (1 - e.ext[2, 1, 1]), abs=1e-15)
    assert world.viewer.get(f"/api/projects/{pid}/export/curves?scenario_id=999999").status_code == 404


def test_all_cohorts_results_workbook_and_bundle(world):
    pid, sid = world.pid, world.sid
    r = world.viewer.get(f"/api/projects/{pid}/export/results?scenario_id={sid}")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames == ["Summary", "Results_by_TermStep", "LGD_to_Target", "Tail_fit", "Parameters", "Warnings"]
    e = _engine_44()
    rows = [[c.value for c in row] for row in wb["Summary"].iter_rows(min_row=5, max_col=15)]
    assert [r_[0] for r_ in rows] == ["VB44", "VBALL"]
    assert rows[0][3] == "Log-normal" and rows[0][14] == pytest.approx(e.averages["lgd_selected"], abs=1e-12)
    ws = wb["Results_by_TermStep"]
    heads = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
    assert heads[:4] == ["Zip", "Category", "TermStep", "Exposure at ts (R)"] and "LGD – log-normal" in heads
    long = [[c.value for c in row] for row in ws.iter_rows(min_row=4, max_col=3)]
    assert sum(1 for r_ in long if r_[0] == "VB44") == 97 and sum(1 for r_ in long if r_[0] == "VBALL") == 329
    first = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    assert first[heads.index("LGD – SELECTED")] == pytest.approx(e.results["lgd_selected"][0], abs=1e-12)
    assert wb["LGD_to_Target"].max_row == 3 + 2 * 360
    params = {row[0].value: row[1].value for row in wb["Parameters"].iter_rows(min_row=3, max_col=2)}
    assert params["Parameter"] == "VB44" and params["MaxBucket"] == 480
    z = world.viewer.get(f"/api/projects/{pid}/export/results?scenario_id={sid}&bundle=zip")
    assert z.status_code == 200 and z.headers["content-type"] == "application/zip"
    import zipfile
    names = sorted(zipfile.ZipFile(io.BytesIO(z.content)).namelist())
    assert names == ["VB44_Base_values.xlsx", "VBALL_Base_values.xlsx"]
    assert world.viewer.get(f"/api/projects/{pid}/export/results?scenario_id={sid}&bundle=tar").status_code == 422
    assert world.viewer.get(f"/api/projects/{pid}/export/results?scenario_id=999999").status_code == 404


def test_outsider_cannot_reach_results_or_exports(world):
    base = f"/api/projects/{world.pid}"
    for url in (f"{base}/matrix", f"{base}/results/{world.sid}/{world.d44}",
                f"{base}/results/{world.sid}/{world.d44}/export?kind=values", f"{base}/export/summary",
                f"{base}/scenarios", f"{base}/curves"):
        assert world.outsider.get(url).status_code == 404, url


# ------------------------------------------------------------ vintages
def test_scenario_on_the_last_ten_years_runs_per_zip(world):
    pid = world.pid
    s = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Ten years",
                          "params": {"method": 3, "target_ts": 360, "max_bucket": 480, "vintage_years": 10}}, headers=H).json()
    out = {o["dataset"]: o for o in world.editor.post(f"/api/projects/{pid}/scenarios/{s['id']}/run", json={}, headers=H).json()}
    assert out["VB44"]["status"] == "ok" and out["VBALL"]["status"] == "ok"
    assert out["VB44"]["summary"]["vintage_filter"] is False            # VB44's first vintage is within ten years
    assert out["VBALL"]["summary"]["vintage_filter"] is True and out["VBALL"]["summary"]["vintages"].startswith("from 2016-08")
    r = world.viewer.get(f"/api/projects/{pid}/results/{s['id']}/{world.dall}").json()
    cfg = r["payload"]["config"]
    assert cfg["vintage_start_effective"] == "2016-08" and cfg["lgd_file_label"] == "LGD (vintages from 2016-08)"
    assert cfg["cohorts_included"] < cfg["cohorts_total"] and r["payload"]["warnings"][0].startswith("Vintages from 2016-08")
    # the exports carry the window
    csv = world.viewer.get(f"/api/projects/{pid}/results/{s['id']}/{world.dall}/export?kind=csv&table=results").text
    assert "LGD (vintages from 2016-08)" in csv.splitlines()[0]
    v = world.viewer.get(f"/api/projects/{pid}/results/{s['id']}/{world.dall}/export?kind=values")
    assert v.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(v.content), read_only=True)
    assert [c.value for c in next(wb["Results"].iter_rows(min_row=1, max_row=1))][2] == "LGD (vintages from 2016-08)"
    f = world.viewer.get(f"/api/projects/{pid}/results/{s['id']}/{world.dall}/export?kind=formula")
    assert f.status_code == 200, f.text
    wf = openpyxl.load_workbook(io.BytesIO(f.content), read_only=True, data_only=True)
    raw_note = wf["Raw_Debug"]["A1"].value
    assert "rebuilt from runoff_triangle.csv" in raw_note and "2016-08" in raw_note
    cfg_rows = {row[0].value: row[1].value for row in wf["Config"].iter_rows(min_row=4, max_row=45, max_col=2) if row[0].value}
    assert cfg_rows["Vintages"].startswith("From 2016-08")
    assert cfg_rows["Last observed bucket in file (max BucketIndex with exposure > 0)"] == cfg["last_obs_file"] < 328
    summary = world.viewer.get(f"/api/projects/{pid}/export/summary")
    ws = openpyxl.load_workbook(io.BytesIO(summary.content), read_only=True)["Summary"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=5, max_col=6)]
    assert any(r[2] == "Ten years" and r[0] == "VBALL" and str(r[5]).startswith("From 2016-08") for r in rows)
    world.editor.delete(f"/api/projects/{pid}/scenarios/{s['id']}", headers=H)


def test_reupload_of_a_zip_stored_without_runoff_updates_it_in_place(world):
    pid, d44 = world.pid, world.d44
    store, cache = world.app.state.store, world.app.state.cache
    with world.app.state.session_factory() as db:
        ds = db.get(Dataset, d44)
        # rewrite the stored data as an old blob: no runoff, as uploads before 9 October 2026 were
        data = RecoveryData.from_bytes(store.get(ds.blob_key))
        old = RecoveryData(category=data.category, meta=data.meta, event_types=data.event_types, raw=data.raw,
                           identical_events=data.identical_events)
        store.put(ds.blob_key, old.to_bytes())
        cache.drop(ds.blob_key)
        ds.profile = old.profile()
        db.commit()
    assert world.viewer.get(f"/api/projects/{pid}").json()["datasets"][0]["profile"]["has_runoff"] is False
    # a filtered run now fails with the re-upload message
    s = world.editor.post(f"/api/projects/{pid}/scenarios", json={"name": "Needs runoff", "params": {"vintage_years": 5}}, headers=H).json()
    out = {o["dataset"]: o for o in world.editor.post(f"/api/projects/{pid}/scenarios/{s['id']}/run", json={"dataset_id": d44}, headers=H).json()}
    assert out["VB44"]["status"] == "error" and "Upload the zip again" in out["VB44"]["error"]
    # the same zip again: updated in place, same id, results out of date
    r = upload(world.editor, pid, ["44"]).json()[0]
    assert r["ok"] is True and r.get("updated") is True and r["message"] == "Updated with vintage data"
    assert r["dataset"]["id"] == d44 and r["dataset"]["profile"]["has_runoff"] is True
    assert world.editor.get(f"/api/projects/{pid}/results/{world.sid}/{d44}").json()["stale"] is True
    out = {o["dataset"]: o for o in world.editor.post(f"/api/projects/{pid}/scenarios/{s['id']}/run", json={"dataset_id": d44}, headers=H).json()}
    assert out["VB44"]["status"] == "ok" and out["VB44"]["summary"]["vintage_filter"] is True
    # and a second re-upload is refused as a duplicate again
    r = upload(world.editor, pid, ["44"]).json()[0]
    assert r["ok"] is False and "already in the project" in r["error"]
    world.editor.delete(f"/api/projects/{pid}/scenarios/{s['id']}", headers=H)
    world.editor.post(f"/api/projects/{pid}/scenarios/{world.sid}/run", json={}, headers=H)


def test_stored_method_3_from_the_reference_curve_days_moves_to_log_normal(world):
    pid, d44 = world.pid, world.d44
    with world.app.state.session_factory() as db:
        s = Scenario(project_id=pid, name="Old method 3", params={"method": 3, "client_cohort": "44", "target_ts": 300,
                                                                  "max_bucket": 420}, created_by=1)
        db.add(s)
        db.flush()
        db.add(ScenarioOverride(scenario_id=s.id, dataset_id=d44, params={"client_cohort": "22", "window": 6}))
        db.add(Result(scenario_id=s.id, dataset_id=d44, status="ok", stale=False, computed_by=1,
                      effective_params={"method": 3, "client_cohort": "22", "window": 6},
                      summary={"lgd_selected": 0.5, "lgd_client": 0.5, "method_label": "Reference curve shape"},
                      payload={"averages": {"lgd_client": 0.5, "lgd_selected": 0.5}, "config": {"method": 3, "client_cohort": "22"},
                               "lgd_ts": {"ts": [1], "lgd_final": [0.5]}, "warnings": []}))
        db.commit()
        sid = s.id
    # before the migration the stored result is flagged as legacy wherever it is shown
    cell = next(c for c in world.viewer.get(f"/api/projects/{pid}/matrix").json()["cells"] if c["scenario_id"] == sid)
    assert cell["legacy"] is True
    assert world.viewer.get(f"/api/projects/{pid}/results/{sid}/{d44}/export?kind=values").status_code == 400
    out = migrate_three_methods(world.app.state.session_factory)
    assert out["stripped"] == 2 and out["stale"] == 1
    assert migrate_three_methods(world.app.state.session_factory) == {"stripped": 0, "stale": 0}
    got = world.viewer.get(f"/api/projects/{pid}/scenarios/{sid}").json()
    assert got["params"]["method"] == 3 and "client_cohort" not in got["params"]
    assert got["overrides"][str(d44)] == {"window": 6}
    r = world.viewer.get(f"/api/projects/{pid}/results/{sid}/{d44}").json()
    assert r["stale"] is True and r["legacy"] is True
    assert world.viewer.get(f"/api/projects/{pid}").json()["legacy_method3"] == 1
    # the stored parameters still load, and a run gives log-normal figures
    out = world.editor.post(f"/api/projects/{pid}/scenarios/{sid}/run", json={"dataset_id": d44}, headers=H).json()[0]
    assert out["status"] == "ok" and out["summary"]["method_label"] == "Log-normal" and out["legacy"] is False
    r = world.viewer.get(f"/api/projects/{pid}/results/{sid}/{d44}").json()
    assert r["effective_params"]["window"] == 6 and "client_cohort" not in r["effective_params"]
    assert r["payload"]["averages"]["lgd_logn"] == r["payload"]["averages"]["lgd_selected"]
    assert world.viewer.get(f"/api/projects/{pid}").json()["legacy_method3"] == 0
    world.editor.delete(f"/api/projects/{pid}/scenarios/{sid}", headers=H)


# ----------------------------------------------------------------- cleanup
def test_delete_zip_removes_its_stored_data_and_results(world):
    pid = world.pid
    assert world.editor.delete(f"/api/projects/{pid}/datasets/{world.dall}", headers=H).status_code == 200
    assert len(list((world.tmp / "blobs" / f"p{pid}").glob("*.npz"))) == 1
    assert world.editor.get(f"/api/projects/{pid}/results/{world.sid}/{world.dall}").status_code == 404
    s = world.editor.get(f"/api/projects/{pid}/scenarios/{world.sid2}").json()
    assert s["overrides"] == {}


def test_password_change_signs_out_other_sessions(world):
    other = TestClient(world.app)
    login(other, VIEWER)
    assert other.get("/api/auth/me").status_code == 200
    r = world.viewer.post("/api/auth/password", json={"current": "wrong", "new": "brand-new-password"}, headers=H)
    assert r.status_code == 400
    r = world.viewer.post("/api/auth/password", json={"current": VIEWER[1], "new": "brand-new-password"}, headers=H)
    assert r.status_code == 200
    assert world.viewer.get("/api/auth/me").status_code == 200      # this browser stays signed in
    assert other.get("/api/auth/me").status_code == 401              # the old session is dead


def test_sign_out_kills_a_copied_cookie(world):
    client = TestClient(world.app)
    login(client, EDITOR)
    stolen = TestClient(world.app)
    stolen.cookies.set("hazard_session", client.cookies.get("hazard_session"))
    assert stolen.get("/api/auth/me").status_code == 200
    assert client.post("/api/auth/logout", headers=H).status_code == 200
    assert stolen.get("/api/auth/me").status_code == 401
    assert world.editor.get("/api/auth/me").status_code == 401       # signed out everywhere
    login(world.editor, EDITOR)


def test_anonymous_upload_is_refused_before_the_body_is_read(world):
    r = upload(world.anon, world.pid, [("x.zip", b"abc")])
    assert r.status_code == 401


def test_api_reference_is_not_public(world):
    assert world.anon.get("/api/openapi.json").status_code == 404
    assert world.anon.get("/api/docs").status_code == 404


def test_deactivated_user_is_signed_out(world):
    users = {u["email"]: u for u in world.admin.get("/api/admin/users").json()}
    world.admin.patch(f"/api/admin/users/{users[OUTSIDER[0]]['id']}", json={"is_active": False}, headers=H)
    assert world.outsider.get("/api/auth/me").status_code == 401
    r = world.outsider.post("/api/auth/login", json={"email": OUTSIDER[0], "password": OUTSIDER[1]}, headers=H)
    assert r.status_code == 401


def test_owner_deletes_project(world):
    assert world.viewer.delete(f"/api/projects/{world.pid}", headers=H).status_code == 403
    assert world.editor.delete(f"/api/projects/{world.pid}", headers=H).status_code == 200
    assert world.editor.get("/api/projects").json() == []
    assert not list((world.tmp / "blobs" / f"p{world.pid}").glob("*.npz"))
