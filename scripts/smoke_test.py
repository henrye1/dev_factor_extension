"""End-to-end check of a running app through its API, using the administrator in .env.

Signs in, creates a temporary project, uploads two zips, runs a scenario, compares the result
with the engine run directly, downloads every export, then deletes the temporary project.
Nothing is left behind except the administrator account.

Usage:  python scripts/smoke_test.py http://127.0.0.1:8000
"""
from __future__ import annotations

import io
import re
import sys
import time
from pathlib import Path

import httpx
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from hazard_ext.engine.core import compute            # noqa: E402
from hazard_ext.engine.curves import builtin_curves   # noqa: E402
from hazard_ext.engine.params import Params            # noqa: E402
from hazard_ext.engine.parse import parse_zip          # noqa: E402

H = {"X-Requested-With": "hazard-ext"}


def env(key: str) -> str:
    for line in (ROOT / ".env").read_text(encoding="utf-8").splitlines():
        m = re.match(rf"^{key}=(.*)$", line.strip())
        if m:
            return m.group(1)
    raise SystemExit(f"{key} is not in .env")


def main(base: str) -> int:
    c = httpx.Client(base_url=base, timeout=600)
    t0 = time.time()

    def step(label, fn):
        t = time.time()
        out = fn()
        print(f"ok   {label}  ({time.time() - t:.1f}s)")
        return out

    def check(resp, code=200):
        if resp.status_code != code:
            raise SystemExit(f"FAIL {resp.request.method} {resp.request.url.path}: {resp.status_code} {resp.text[:400]}")
        return resp

    step("health", lambda: check(c.get("/api/health")))
    me = step("sign in as the administrator in .env", lambda: check(c.post(
        "/api/auth/login", json={"email": env("ADMIN_EMAIL"), "password": env("ADMIN_PASSWORD")}, headers=H)).json())
    assert me["is_admin"]

    pid = step("create a temporary project", lambda: check(c.post(
        "/api/projects", json={"name": f"Smoke test {int(time.time())}"}, headers=H), 201).json()["id"])
    try:
        ids = {}
        for name in ("44", "22"):
            path = ROOT / f"debug ({name}).zip"
            out = step(f"upload debug ({name}).zip ({path.stat().st_size / 1e6:.0f} MB)", lambda: check(c.post(
                f"/api/projects/{pid}/datasets", files=[("files", (path.name, path.read_bytes(), "application/zip"))], headers=H)).json())
            assert out[0]["ok"], out[0]
            ids[name] = out[0]["dataset"]["id"]

        sid = step("create a scenario", lambda: check(c.post(
            f"/api/projects/{pid}/scenarios", json={"name": "Base", "params": {"target_ts": 360, "max_bucket": 480}}, headers=H), 201).json()["id"])
        step("override for one zip", lambda: check(c.put(
            f"/api/projects/{pid}/scenarios/{sid}/overrides/{ids['22']}", json={"params": {"method": 2, "last_ts": 120}}, headers=H)))
        runs = step("run the scenario on both zips", lambda: check(c.post(
            f"/api/projects/{pid}/scenarios/{sid}/run", json={}, headers=H)).json())
        assert all(r["status"] == "ok" for r in runs), runs

        res = step("read the stored result", lambda: check(c.get(f"/api/projects/{pid}/results/{sid}/{ids['44']}")).json())
        expect = compute(parse_zip(str(ROOT / "debug (44).zip")), Params(target_ts=360, max_bucket=480), builtin_curves()["44"])
        diff = max(abs(a - b) for a, b in zip(res["payload"]["lgd_ts"]["lgd_final"], expect.lgd_ts["lgd_final"]))
        assert diff < 1e-12, diff
        print(f"ok   stored LGD equals the engine run directly (max difference {diff:.1e})")
        r22 = check(c.get(f"/api/projects/{pid}/results/{sid}/{ids['22']}")).json()
        assert r22["effective_params"]["method"] == 2 and r22["effective_params"]["last_ts"] == 120

        step("matrix", lambda: check(c.get(f"/api/projects/{pid}/matrix")))
        step("chart data", lambda: check(c.get(f"/api/projects/{pid}/results/{sid}/{ids['44']}/curve?ts=48")))
        base_url = f"/api/projects/{pid}/results/{sid}/{ids['44']}/export"
        step("CSV export", lambda: check(c.get(base_url + "?kind=csv&table=lgd_ts")))
        v = step("Excel values export", lambda: check(c.get(base_url + "?kind=values")))
        assert "Results" in openpyxl.load_workbook(io.BytesIO(v.content), read_only=True).sheetnames
        f = step("Excel formula export", lambda: check(c.get(base_url + "?kind=formula")))
        assert "Raw_Debug" in openpyxl.load_workbook(io.BytesIO(f.content), read_only=True).sheetnames
        step("summary workbook", lambda: check(c.get(f"/api/projects/{pid}/export/summary")))

        s = check(c.get(f"/api/projects/{pid}/scenarios/{sid}")).json()
        step("edit the scenario (marks results out of date)", lambda: check(c.put(
            f"/api/projects/{pid}/scenarios/{sid}", json={"params": {**s["params"], "window": 18}}, headers=H)))
        assert check(c.get(f"/api/projects/{pid}/results/{sid}/{ids['44']}")).json()["stale"] is True
        assert c.get(base_url + "?kind=values").status_code == 409
    finally:
        step("delete the temporary project", lambda: check(c.delete(f"/api/projects/{pid}", headers=H)))
    assert all(p["id"] != pid for p in check(c.get("/api/projects")).json())
    print(f"All checks passed in {time.time() - t0:.0f}s")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8000"))
